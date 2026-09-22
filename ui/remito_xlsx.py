"""Excel del pedido de equipos, con el formato que se usaba en AsistenteCaja: IMEI, MODELO, DESTINO y PEDIDO."""

HEADERS = ["IMEI", "MODELO", "DESTINO ", "PEDIDO"]
HEADERS_CHIPS = ["TIPO DE SIM", "PRIMER CHIP", "ÚLTIMO CHIP", "CANTIDAD", "DESTINO", "PEDIDO"]


def _libro(headers, anchos):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, Side

    wb = Workbook()
    ws = wb.active
    ws.title = "Hoja1"
    for letra, ancho in anchos.items():
        ws.column_dimensions[letra].width = ancho
    ws.row_dimensions[1].height = 29.25
    borde = Border(*(Side(style="thin"),) * 4)
    centrado = Alignment(horizontal="center", vertical="center")
    for col, texto in enumerate(headers, start=1):
        c = ws.cell(1, col, texto)
        c.font, c.alignment, c.border = Font(name="Arial", size=16, bold=True), centrado, borde
    return wb, ws, borde, centrado


def generar(path, destino, pedido, items):
    """`items`: [(imei, modelo)]. Necesita openpyxl (ImportError si no está instalado)."""
    from openpyxl.styles import Font, PatternFill

    wb, ws, borde, centrado = _libro(HEADERS, {"A": 30.71, "B": 34.43, "C": 20.71, "D": 20.71})
    blanco = PatternFill(fill_type="solid", fgColor="FFFFFFFF")
    pedido_valor = int(pedido) if str(pedido).strip().isdigit() else pedido
    for fila, (imei, modelo) in enumerate(items, start=2):
        c = ws.cell(fila, 1, str(imei))   # texto, para no perder dígitos por notación científica
        c.font, c.alignment, c.border, c.fill, c.number_format = Font(name="Arial", size=10), centrado, borde, blanco, "@"
        c = ws.cell(fila, 2, modelo)
        c.font, c.border = Font(name="Arial", size=11, color="FF333333"), borde
        c = ws.cell(fila, 3, destino)
        c.font, c.alignment, c.border, c.fill = Font(name="Arial", size=10), centrado, borde, blanco
        c = ws.cell(fila, 4, pedido_valor)
        c.font, c.alignment, c.border = Font(name="Arial", size=10, color="FF000000"), centrado, borde
    ws.auto_filter.ref = f"A1:D{1 + len(items)}"
    wb.save(path)


def generar_chips(path, destino, pedido, items):
    """`items`: [(tipo de SIM, primer chip, último chip, cantidad)], para mandarle a administración el rango de
    cada caja (no se guarda en la base: Nodo solo lleva la cantidad). Necesita openpyxl."""
    from openpyxl.styles import Font

    wb, ws, borde, centrado = _libro(HEADERS_CHIPS, {"A": 34.43, "B": 20.71, "C": 20.71, "D": 14.29, "E": 20.71,
                                                     "F": 20.71})
    pedido_valor = int(pedido) if str(pedido).strip().isdigit() else pedido
    for fila, (tipo, primero, ultimo, cantidad) in enumerate(items, start=2):
        c = ws.cell(fila, 1, tipo)
        c.font, c.border = Font(name="Arial", size=11, color="FF333333"), borde
        for col, valor in ((2, primero), (3, ultimo)):
            c = ws.cell(fila, col, valor)
            c.font, c.alignment, c.border, c.number_format = Font(name="Arial", size=10), centrado, borde, "@"
        c = ws.cell(fila, 4, cantidad)
        c.font, c.alignment, c.border = Font(name="Arial", size=10), centrado, borde
        c = ws.cell(fila, 5, destino)
        c.font, c.alignment, c.border = Font(name="Arial", size=10), centrado, borde
        c = ws.cell(fila, 6, pedido_valor)
        c.font, c.alignment, c.border = Font(name="Arial", size=10, color="FF000000"), centrado, borde
    ws.auto_filter.ref = f"A1:F{1 + len(items)}"
    wb.save(path)
