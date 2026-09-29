"""Excel de las gestiones: las filas que se ven en la tabla (filtro, período y Buscar), con todos sus datos."""
from datetime import datetime

from ui.formatting import fmt_money

FORMATO_FECHA = "DD/MM/YYYY"
FORMATO_FECHA_HORA = "DD/MM/YYYY HH:MM"
FORMATO_MONTO = '"$" #,##0.00'


def columnas(screen):
    """Primero las columnas de la tabla, en su orden; después el resto de los campos (los que solo están en el
    formulario), en el orden del formulario."""
    vistas = list(screen.columns)
    return vistas + [f for f in screen.fields if f not in vistas]


def _valor(screen, field, row):
    """El valor de la celda: números y fechas como tales (para poder sumar y filtrar en Excel); el resto, como se ve."""
    value = row.get(field.key)
    if field.kind in ("money", "int"):
        return value or 0
    if field.kind in ("date", "datetime"):
        if not value:
            return None
        return datetime.strptime(value[:16], "%Y-%m-%d %H:%M") if len(value) > 10 else datetime.strptime(value, "%Y-%m-%d")
    if field.kind == "bool":
        return "Sí" if value else "No"
    if field.kind == "pagos":
        return "; ".join(f"{codigo} {fmt_money(monto)}" for _, monto, codigo in value or [])
    return screen._display(field, row)


def generar(path, screen, rows):
    """Guarda `rows` (filas de la pantalla de gestiones `screen`) en un Excel. Necesita openpyxl."""
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    campos = columnas(screen)
    wb = Workbook()
    ws = wb.active
    ws.title = screen.title.split("·")[-1].strip()[:31]
    ws.append([f.label for f in campos])
    for c in ws[1]:
        c.font = Font(bold=True)
    for row in rows:
        ws.append([_valor(screen, f, row) for f in campos])
    for col, f in enumerate(campos, start=1):
        letra = get_column_letter(col)
        ws.column_dimensions[letra].width = min(40, max(12, len(f.label) + 2, f.width // 7))
        formato = {"money": FORMATO_MONTO, "date": FORMATO_FECHA, "datetime": FORMATO_FECHA_HORA}.get(f.kind)
        if formato:
            for (c,) in ws.iter_rows(min_row=2, min_col=col, max_col=col):
                c.number_format = formato
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(campos))}{1 + len(rows)}"
    wb.save(path)
